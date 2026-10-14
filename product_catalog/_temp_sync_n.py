import csv,json,os,re
from collections import Counter
from urllib.parse import urlparse
from openpyxl import load_workbook
ROOT=os.path.dirname(__file__);TODAY='2026-10-13';FILE='N-networking-gps-communication.csv'
def read(n):
 with open(os.path.join(ROOT,n),encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def write(n,rs,cols=None):
 with open(os.path.join(ROOT,n),'w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=cols or list(rs[0]),extrasaction='ignore');w.writeheader();w.writerows(rs)
rows=read(FILE);fields=list(rows[0]);counts=Counter(r['subcategory'] for r in rows)
files=[f'{c}-{n}.csv' for c,n in [('A','mobile-phones-tablets'),('B','computers-laptops-it-equipment'),('C','tvs-video-home-entertainment'),('D','audio-headphones-speakers'),('E','large-home-appliances'),('F','kitchen-appliances'),('G','cleaning-laundry-floor-care'),('H','air-conditioning-cooling-air-quality'),('I','personal-care-grooming-appliances'),('J','cameras-photography-imaging'),('K','gaming-virtual-reality'),('L','wearables-smart-devices'),('M','smart-home-security-surveillance'),('N','networking-gps-communication')]]
allrows=[]
for name in files:allrows.extend(read(name))
write('all_products.csv',allrows,fields)
plan=read('category_plan.csv')
for p in plan:
 if p['department_code']=='N':p['status']='COMPLETE'
 elif p['department_code']=='O':p['status']='IN_PROGRESS'
write('category_plan.csv',plan)
summary=[r for r in read('master_summary.csv') if r['category_code'] not in ('N','O')]
for sub,n in [('ALL',len(rows))]+sorted(counts.items()):summary.append(dict(category_code='N',department='Networking, GPS & Communication',subcategory=sub,products_count=str(n),processing_status='COMPLETE',csv_file=FILE))
summary.append(dict(category_code='O',department='Printers, Scanners & Office Electronics',subcategory='ALL',products_count='0',processing_status='IN_PROGRESS',csv_file='O-printers-scanners-office-electronics.csv'))
write('master_summary.csv',summary)
domains=Counter(re.sub(r'^www\.','',urlparse(r['product_source_url']).hostname or '') for r in rows)
log=[r for r in read('collection_log.csv') if r['category_code'] not in ('N','O')]
log.append({'category_code':'N','product_type':'ALL','status':'COMPLETE','product_count':str(len(rows)),'brand_count':str(len(set(r['brand'] for r in rows))),'source_domain':'|'.join(sorted(domains)),'rows_attributed_to_domain':json.dumps(dict(sorted(domains.items())),separators=(',',':')),'last_updated':TODAY,'notes':f'All {len(counts)} supplied subtypes represented; counts {dict(sorted(counts.items()))}. Release year missing {sum(not r["release_year"] for r in rows)}; price missing {sum(not r["official_price"] for r in rows)}; image missing {sum(not r["image_url"] for r in rows)}; empty specs JSON {sum(r["specifications_json"]=="{}" for r in rows)}. Missing/uncertain fields are flagged per row.'})
for sub,n in sorted(counts.items()):log.append({'category_code':'N','product_type':sub,'status':'COMPLETE','product_count':str(n),'brand_count':str(len(set(r['brand'] for r in rows if r['subcategory']==sub))),'source_domain':'','rows_attributed_to_domain':'','last_updated':TODAY,'notes':'Official manufacturer URLs are included per row; verify regional fit, release dates, prices, and image URLs before import.'})
log.append({'category_code':'O','product_type':'ALL','status':'IN_PROGRESS','product_count':'0','brand_count':'0','source_domain':'','rows_attributed_to_domain':'','last_updated':TODAY,'notes':'Next active category: Printers, Scanners & Office Electronics.'})
write('collection_log.csv',log)
wb=load_workbook(os.path.join(ROOT,'master_catalogue.xlsx'))
for n in ['N','N - Stats','Category Summary']:
 if n not in wb.sheetnames:wb.create_sheet(n)
ws=wb['N'];ws.delete_rows(1,ws.max_row);ws.append(fields)
for r in rows:ws.append([r.get(k,'') for k in fields])
st=wb['N - Stats'];st.delete_rows(1,st.max_row);st.append(['N Product Statistics',None]);st.append(['Metric','Count'])
for k,v in [('Total products',len(rows)),('Distinct brands',len(set(r['brand'] for r in rows))),('Release year verified',sum(bool(r['release_year']) for r in rows)),('Official price listed',sum(bool(r['official_price']) for r in rows)),('Image URL present',sum(bool(r['image_url']) for r in rows)),('Release year unknown',sum(not r['release_year'] for r in rows)),('Price unavailable',sum(not r['official_price'] for r in rows)),('Image URL missing',sum(not r['image_url'] for r in rows)),('Detailed specs unavailable',sum(r['specifications_json']=='{}' for r in rows))]:st.append([k,v])
st.append([None,None]);st.append(['Subtype','Products'])
for sub,n in sorted(counts.items()):st.append([sub,n])
cs=wb['Category Summary'];cs.delete_rows(1,cs.max_row);cs.append(['Category','Department','Subtype','Products','Status','CSV file'])
for r in sorted(summary,key=lambda x:(x['category_code'],x['subcategory'])):cs.append([r['category_code'],r['department'],r['subcategory'],int(r['products_count']),r['processing_status'],r['csv_file']])
wb.save(os.path.join(ROOT,'master_catalogue.xlsx'))
print('N validated:',len(rows),'products,',len(set(r['brand'] for r in rows)),'brands,',len(counts),'subtypes; duplicate keys',len(rows)-len(set((r['brand'].lower(),r['model'].lower()) for r in rows)))
print('Combined rows',len(allrows),'; O is now sole active category.')
