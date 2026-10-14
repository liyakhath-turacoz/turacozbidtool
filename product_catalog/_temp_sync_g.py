import csv,json,os,re
from collections import Counter
from urllib.parse import urlparse
from openpyxl import load_workbook
ROOT=os.path.dirname(__file__);TODAY='2026-10-13';FILE='G-cleaning-laundry-floor-care.csv';FIELDS=list(next(csv.DictReader(open(os.path.join(ROOT,FILE),encoding='utf-8-sig'))).keys())
def read(n):
 with open(os.path.join(ROOT,n),encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def write(n,rs,cols=None):
 with open(os.path.join(ROOT,n),'w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=cols or list(rs[0]),extrasaction='ignore');w.writeheader();w.writerows(rs)
rows=read(FILE);counts=Counter(r['subcategory'] for r in rows)
allrows=[]
for n in ['A-mobile-phones-tablets.csv','B-computers-laptops-it-equipment.csv','C-tvs-video-home-entertainment.csv','D-audio-headphones-speakers.csv','E-large-home-appliances.csv','F-kitchen-appliances.csv',FILE]:allrows.extend(read(n))
write('all_products.csv',allrows,FIELDS)
plan=read('category_plan.csv')
for p in plan:
 if p['department_code']=='G':p['status']='COMPLETE'
 elif p['department_code']=='H':p['status']='IN_PROGRESS'
write('category_plan.csv',plan)
summary=read('master_summary.csv')
for r in summary:
 if r['category_code']=='H':r['processing_status']='IN_PROGRESS'
summary=[r for r in summary if r['category_code']!='G']
summary.append(dict(category_code='G',department='Cleaning, Laundry & Floor Care',subcategory='ALL',products_count=str(len(rows)),processing_status='COMPLETE',csv_file=FILE))
for sub,n in sorted(counts.items()):summary.append(dict(category_code='G',department='Cleaning, Laundry & Floor Care',subcategory=sub,products_count=str(n),processing_status='COMPLETE',csv_file=FILE))
write('master_summary.csv',summary)
domains=Counter(re.sub(r'^www\.','',urlparse(r['product_source_url']).hostname or '') for r in rows)
log=[r for r in read('collection_log.csv') if r['category_code'] not in ('G','H')]
log.append({'category_code':'G','product_type':'ALL','status':'COMPLETE','product_count':str(len(rows)),'brand_count':str(len(set(r['brand'] for r in rows))),'source_domain':'|'.join(sorted(domains)),'rows_attributed_to_domain':json.dumps(dict(sorted(domains.items())),separators=(',',':')),'last_updated':TODAY,'notes':f'All 6 supplied subtypes represented; subtype counts {dict(sorted(counts.items()))}. Release year missing {sum(not r["release_year"] for r in rows)}; official price missing {sum(not r["official_price"] for r in rows)}; image URL missing {sum(not r["image_url"] for r in rows)}; detailed specification JSON missing {sum(r["specifications_json"]=="{}" for r in rows)}. Missing fields left blank and flagged.'})
for sub,n in sorted(counts.items()):log.append({'category_code':'G','product_type':sub,'status':'COMPLETE','product_count':str(n),'brand_count':str(len(set(r['brand'] for r in rows if r['subcategory']==sub))),'source_domain':'','rows_attributed_to_domain':'','last_updated':TODAY,'notes':'Official manufacturer catalogue/product pages recorded per row; age/specification gaps are explicit.'})
log.append({'category_code':'H','product_type':'ALL','status':'IN_PROGRESS','product_count':'0','brand_count':'0','source_domain':'','rows_attributed_to_domain':'','last_updated':TODAY,'notes':'Current category; begin subtype-by-subtype collection after G export validation.'})
write('collection_log.csv',log)
wb=load_workbook(os.path.join(ROOT,'master_catalogue.xlsx'))
for nm in ['G','G - Stats','Category Summary']:
 if nm not in wb.sheetnames:wb.create_sheet(nm)
ws=wb['G'];ws.delete_rows(1,ws.max_row);ws.append(FIELDS)
for r in rows:ws.append([r.get(k,'') for k in FIELDS])
st=wb['G - Stats'];st.delete_rows(1,st.max_row);st.append(['G Product Statistics',None]);st.append(['Metric','Count'])
for k,v in [('Total products',len(rows)),('Distinct brands',len(set(r['brand'] for r in rows))),('Release year verified',sum(bool(r['release_year']) for r in rows)),('Official price listed',sum(bool(r['official_price']) for r in rows)),('Image URL present',sum(bool(r['image_url']) for r in rows)),('Release year unknown',sum(not r['release_year'] for r in rows)),('Price unavailable',sum(not r['official_price'] for r in rows)),('Image URL missing',sum(not r['image_url'] for r in rows)),('Detailed specs unavailable',sum(r['specifications_json']=='{}' for r in rows))]:st.append([k,v])
st.append([None,None]);st.append(['Subtype','Products'])
for sub,n in sorted(counts.items()):st.append([sub,n])
cs=wb['Category Summary'];cs.delete_rows(1,cs.max_row);cs.append(['Category','Department','Subtype','Products','Status','CSV file'])
for r in sorted(summary,key=lambda x:(x['category_code'],x['subcategory'])):cs.append([r['category_code'],r['department'],r['subcategory'],int(r['products_count']),r['processing_status'],r['csv_file']])
wb.save(os.path.join(ROOT,'master_catalogue.xlsx'))
keys=[(r['brand'].lower(),r['model'].lower()) for r in rows]
print('G validation:',len(rows),'products;',len(set(r['brand'] for r in rows)),'brands;',len(counts),'subtypes; duplicate keys',len(keys)-len(set(keys)))
print('G gaps: year',sum(not r['release_year'] for r in rows),'price',sum(not r['official_price'] for r in rows),'image',sum(not r['image_url'] for r in rows),'specs',sum(r['specifications_json']=='{}' for r in rows),'blank sources',sum(not r['product_source_url'] for r in rows))
print('Combined rows',len(allrows),'; H is now sole active category.')
