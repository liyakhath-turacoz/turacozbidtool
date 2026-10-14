import csv, json, os, re
from collections import Counter
from urllib.parse import urlparse
from openpyxl import load_workbook

ROOT=os.path.dirname(__file__); TODAY='2026-10-13'; FILE='J-cameras-photography-imaging.csv'
def read(name):
    with open(os.path.join(ROOT,name),encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def write(name,rows,fields=None):
    with open(os.path.join(ROOT,name),'w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields or list(rows[0]),extrasaction='ignore');w.writeheader();w.writerows(rows)
rows=read(FILE); fields=list(rows[0]); counts=Counter(r['subcategory'] for r in rows)
allrows=[]
for code,name in [('A','A-mobile-phones-tablets.csv'),('B','B-computers-laptops-it-equipment.csv'),('C','C-tvs-video-home-entertainment.csv'),('D','D-audio-headphones-speakers.csv'),('E','E-large-home-appliances.csv'),('F','F-kitchen-appliances.csv'),('G','G-cleaning-laundry-floor-care.csv'),('H','H-air-conditioning-cooling-air-quality.csv'),('I','I-personal-care-grooming-appliances.csv'),('J',FILE)]:allrows.extend(read(name))
write('all_products.csv',allrows,fields)
plan=read('category_plan.csv')
for p in plan:
    if p['department_code']=='J':p['status']='COMPLETE'
    elif p['department_code']=='K':p['status']='IN_PROGRESS'
write('category_plan.csv',plan)
summary=[r for r in read('master_summary.csv') if r['category_code']!='J']
for sub,n in [('ALL',len(rows))]+sorted(counts.items()):summary.append(dict(category_code='J',department='Cameras, Photography & Imaging',subcategory=sub,products_count=str(n),processing_status='COMPLETE',csv_file=FILE))
summary.append(dict(category_code='K',department='Gaming & Virtual Reality',subcategory='ALL',products_count='0',processing_status='IN_PROGRESS',csv_file='K-gaming-virtual-reality.csv'))
write('master_summary.csv',summary)
domains=Counter(re.sub(r'^www\.','',urlparse(r['product_source_url']).hostname or '') for r in rows)
log=[r for r in read('collection_log.csv') if r['category_code'] not in ('J','K')]
log.append({'category_code':'J','product_type':'ALL','status':'COMPLETE','product_count':str(len(rows)),'brand_count':str(len(set(r['brand'] for r in rows))),'source_domain':'|'.join(sorted(domains)),'rows_attributed_to_domain':json.dumps(dict(sorted(domains.items())),separators=(',',':')),'last_updated':TODAY,'notes':f'All {len(counts)} represented subtypes; counts {dict(sorted(counts.items()))}. Release year missing {sum(not r["release_year"] for r in rows)}; official price missing {sum(not r["official_price"] for r in rows)}; image missing {sum(not r["image_url"] for r in rows)}; specifications empty {sum(r["specifications_json"]=="{}" for r in rows)}. Missing fields are flagged per record; no values invented.'})
for sub,n in sorted(counts.items()):log.append({'category_code':'J','product_type':sub,'status':'COMPLETE','product_count':str(n),'brand_count':str(len(set(r['brand'] for r in rows if r['subcategory']==sub))),'source_domain':'','rows_attributed_to_domain':'','last_updated':TODAY,'notes':'Manufacturer sources linked in product rows; blank release year, price, image and specifications require review before import.'})
log.append({'category_code':'K','product_type':'ALL','status':'IN_PROGRESS','product_count':'0','brand_count':'0','source_domain':'','rows_attributed_to_domain':'','last_updated':TODAY,'notes':'Next active category: Gaming & Virtual Reality.'})
write('collection_log.csv',log)
wb=load_workbook(os.path.join(ROOT,'master_catalogue.xlsx'))
for name in ['J','J - Stats','Category Summary']:
    if name not in wb.sheetnames:wb.create_sheet(name)
ws=wb['J'];ws.delete_rows(1,ws.max_row);ws.append(fields)
for r in rows:ws.append([r.get(k,'') for k in fields])
st=wb['J - Stats'];st.delete_rows(1,st.max_row);st.append(['J Product Statistics',None]);st.append(['Metric','Count'])
for label,val in [('Total products',len(rows)),('Distinct brands',len(set(r['brand'] for r in rows))),('Release year verified',sum(bool(r['release_year']) for r in rows)),('Official price listed',sum(bool(r['official_price']) for r in rows)),('Image URL present',sum(bool(r['image_url']) for r in rows)),('Release year unknown',sum(not r['release_year'] for r in rows)),('Price unavailable',sum(not r['official_price'] for r in rows)),('Image URL missing',sum(not r['image_url'] for r in rows)),('Detailed specs unavailable',sum(r['specifications_json']=='{}' for r in rows))]:st.append([label,val])
st.append([None,None]);st.append(['Subtype','Products'])
for sub,n in sorted(counts.items()):st.append([sub,n])
cs=wb['Category Summary'];cs.delete_rows(1,cs.max_row);cs.append(['Category','Department','Subtype','Products','Status','CSV file'])
for r in sorted(summary,key=lambda x:(x['category_code'],x['subcategory'])):cs.append([r['category_code'],r['department'],r['subcategory'],int(r['products_count']),r['processing_status'],r['csv_file']])
wb.save(os.path.join(ROOT,'master_catalogue.xlsx'))
print('J validation:',len(rows),'products;',len(set(r['brand'] for r in rows)),'brands;',len(counts),'subtypes; duplicate keys',len(rows)-len(set((r['brand'].lower(),r['model'].lower()) for r in rows)))
print('J gaps: year',sum(not r['release_year'] for r in rows),'price',sum(not r['official_price'] for r in rows),'image',sum(not r['image_url'] for r in rows),'specs',sum(r['specifications_json']=='{}' for r in rows))
print('Combined rows',len(allrows),'; K is now sole active category.')
