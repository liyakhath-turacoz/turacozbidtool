import csv,json,os,re
from collections import Counter
from urllib.parse import urlparse
from openpyxl import load_workbook
ROOT=os.path.dirname(__file__);TODAY='2026-10-13';FILE='K-gaming-virtual-reality.csv'
def read(n):
 with open(os.path.join(ROOT,n),encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def write(n,rs,cols=None):
 with open(os.path.join(ROOT,n),'w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=cols or list(rs[0]),extrasaction='ignore');w.writeheader();w.writerows(rs)
rows=read(FILE);fields=list(rows[0]);counts=Counter(r['subcategory'] for r in rows)
files=['A-mobile-phones-tablets.csv','B-computers-laptops-it-equipment.csv','C-tvs-video-home-entertainment.csv','D-audio-headphones-speakers.csv','E-large-home-appliances.csv','F-kitchen-appliances.csv','G-cleaning-laundry-floor-care.csv','H-air-conditioning-cooling-air-quality.csv','I-personal-care-grooming-appliances.csv','J-cameras-photography-imaging.csv',FILE]
allrows=[]
for name in files:allrows.extend(read(name))
write('all_products.csv',allrows,fields)
plan=read('category_plan.csv')
for p in plan:
 if p['department_code']=='K':p['status']='COMPLETE'
 elif p['department_code']=='L':p['status']='IN_PROGRESS'
write('category_plan.csv',plan)
summary=[r for r in read('master_summary.csv') if r['category_code']!='K']
for sub,n in [('ALL',len(rows))]+sorted(counts.items()):summary.append(dict(category_code='K',department='Gaming & Virtual Reality',subcategory=sub,products_count=str(n),processing_status='COMPLETE',csv_file=FILE))
summary.append(dict(category_code='L',department='Wearables & Smart Devices',subcategory='ALL',products_count='0',processing_status='IN_PROGRESS',csv_file='L-wearables-smart-devices.csv'))
write('master_summary.csv',summary)
domains=Counter(re.sub(r'^www\.','',urlparse(r['product_source_url']).hostname or '') for r in rows)
log=[r for r in read('collection_log.csv') if r['category_code'] not in ('K','L')]
log.append({'category_code':'K','product_type':'ALL','status':'COMPLETE','product_count':str(len(rows)),'brand_count':str(len(set(r['brand'] for r in rows))),'source_domain':'|'.join(sorted(domains)),'rows_attributed_to_domain':json.dumps(dict(sorted(domains.items())),separators=(',',':')),'last_updated':TODAY,'notes':f'All {len(counts)} listed subtypes represented; counts {dict(sorted(counts.items()))}. Year missing {sum(not r["release_year"] for r in rows)}; price missing {sum(not r["official_price"] for r in rows)}; image missing {sum(not r["image_url"] for r in rows)}; empty specification JSON {sum(r["specifications_json"]=="{}" for r in rows)}. Missing data is flagged by record.'})
for sub,n in sorted(counts.items()):log.append({'category_code':'K','product_type':sub,'status':'COMPLETE','product_count':str(n),'brand_count':str(len(set(r['brand'] for r in rows if r['subcategory']==sub))),'source_domain':'','rows_attributed_to_domain':'','last_updated':TODAY,'notes':'Official product or catalogue URLs listed per row; verify regional model, release date and configuration before import.'})
log.append({'category_code':'L','product_type':'ALL','status':'IN_PROGRESS','product_count':'0','brand_count':'0','source_domain':'','rows_attributed_to_domain':'','last_updated':TODAY,'notes':'Next active category: Wearables & Smart Devices.'})
write('collection_log.csv',log)
wb=load_workbook(os.path.join(ROOT,'master_catalogue.xlsx'))
for n in ['K','K - Stats','Category Summary']:
 if n not in wb.sheetnames:wb.create_sheet(n)
ws=wb['K'];ws.delete_rows(1,ws.max_row);ws.append(fields)
for r in rows:ws.append([r.get(k,'') for k in fields])
st=wb['K - Stats'];st.delete_rows(1,st.max_row);st.append(['K Product Statistics',None]);st.append(['Metric','Count'])
for k,v in [('Total products',len(rows)),('Distinct brands',len(set(r['brand'] for r in rows))),('Release year verified',sum(bool(r['release_year']) for r in rows)),('Official price listed',sum(bool(r['official_price']) for r in rows)),('Image URL present',sum(bool(r['image_url']) for r in rows)),('Release year unknown',sum(not r['release_year'] for r in rows)),('Price unavailable',sum(not r['official_price'] for r in rows)),('Image URL missing',sum(not r['image_url'] for r in rows)),('Detailed specs unavailable',sum(r['specifications_json']=='{}' for r in rows))]:st.append([k,v])
st.append([None,None]);st.append(['Subtype','Products'])
for sub,n in sorted(counts.items()):st.append([sub,n])
cs=wb['Category Summary'];cs.delete_rows(1,cs.max_row);cs.append(['Category','Department','Subtype','Products','Status','CSV file'])
for r in sorted(summary,key=lambda x:(x['category_code'],x['subcategory'])):cs.append([r['category_code'],r['department'],r['subcategory'],int(r['products_count']),r['processing_status'],r['csv_file']])
wb.save(os.path.join(ROOT,'master_catalogue.xlsx'))
print('K validation',len(rows),'products,',len(set(r['brand'] for r in rows)),'brands,',len(counts),'subtypes; duplicate keys',len(rows)-len(set((r['brand'].lower(),r['model'].lower()) for r in rows)))
print('Combined rows',len(allrows),'; L is now sole active category.')
