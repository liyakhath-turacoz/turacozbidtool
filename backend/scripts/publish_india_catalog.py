"""Validate and install a staged India catalogue, backing up the old data first.

Dry run by default. --apply replaces the CSV/XLSX source and visible DB catalogue,
then removes unreferenced old products. Bid/cart product references are preserved.
"""
import argparse
from collections import Counter
import csv
from datetime import datetime
import gzip
import hashlib
import json
from pathlib import Path
import shutil
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from backend.catalog_validation import validate_snapshot
from backend.catalog_importer import connect
from backend import catalogue_app_import
from openpyxl import Workbook


def write_csv(path, fields, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fields)
        writer.writeheader()
        writer.writerows(rows)


def documentation(folder, rows, report):
    lines = ['# Current India store catalogue', '',
        f'Verified on **{report["as_of"]}**: **{len(rows):,} purchasable variants**, **{len(report["brands"])} brands**, **{len(report["categories"])} departments**.', '',
        'Every visible row has an India store product listing, positive INR selling price, technical details, a purchase URL and a decoded, locally stored product photograph. Stock is not an inclusion criterion; recheck the current product listing and price before purchase.', '',
        '## Coverage', '', '| Department | Products |', '|---|---:|']
    lines += [f'| {name} | {count} |' for name,count in report['categories'].items()]
    lines += ['', '## Evidence and files', '',
        '- all_products.csv, A–P exports, master_catalogue.xlsx and the root product_inventory.xlsx contain the same replacement records.',
        '- availability_evidence.json records the exact India store listing/variant, observed date, selling price, original photo URL and downloaded image SHA-256.',
        '- verification_report.json records scan totals, rejected inventory and coverage limitations.',
        '- images/ contains the product photos served at /catalog-images/. Web and mobile use the same assets.', '',
        '## Acceptance rules', '',
        f'An India store product page must show a positive INR price and exact variant details. Storefront currency is independently checked. Store listings created before {report["listing_cutoff"]} are excluded. Samsung variants use explicit new-condition Offers from current India product pages instead of a listing-age estimate.', '',
        'Store listing dates are NOT manufacturer release dates. Unknown release dates and years remain blank. Coverage is a verified subset of the configured stores; it does not claim every Indian store or every latest launch.', '',
        'Refurbished, renewed, pre-owned, preorder, gift/bundle, incomplete, unmatched-colour-image and known legacy-generation records are excluded. Exact name and SKU duplicates are removed. Missing metadata is not invented.', '',
        '## Refresh and deployment', '',
        'Install backend/scripts/requirements-catalog.txt. Run rebuild_india_catalog.py with the current India verification date and desired listing cutoff, then inspect .firecrawl/india-catalog-staged/verification_report.json. Run publish_india_catalog.py without flags to validate; --apply publishes and imports.', '',
        'Publication rejects empty/suspiciously small snapshots, product-page evidence older than 30 days, missing details, duplicate identities or mismatched/missing downloaded photos. An imported snapshot does not refresh itself.', '',
        'Backups are private under backend/recovery/ and excluded from Git. Unreferenced old products are deleted; historical products remain hidden for bid/cart references. Legacy import schedules and automatic startup bundle imports are disabled by default.', '',
        'Deploy product_catalog/images with the backend. Restart a backend running without reload to enable the /catalog-images/ route.']
    body='\n'.join(lines)+'\n'
    (folder/'README.md').write_text(body,encoding='utf-8')


def exports(folder, rows, fields, plans):
    write_csv(folder / 'all_products.csv', fields, rows)
    summaries, brands, coverage = [], [], []
    wb = Workbook()
    wb.remove(wb.active)
    for plan in plans:
        code = plan['department_code']
        group = [r for r in rows if r['category_code'] == code]
        filename = next(folder.glob(code + '-*.csv')).name
        write_csv(folder / filename, fields, group)
        sheet = wb.create_sheet(code)
        sheet.append(fields)
        for row in group:
            sheet.append([row.get(key, '') for key in fields])
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        summaries.append({'category_code': code, 'department': plan['department'], 'products': len(group), 'brands': len({r['brand'] for r in group}), 'status': 'VERIFIED_SUBSET' if group else 'NO_VERIFIED_INVENTORY'})
        coverage.append({'category_code':code, 'products':len(group), 'with_price':len(group), 'with_image':len(group), 'with_specs':len(group), 'with_store_source':len(group), 'with_verified_release_date':0})
        brands.extend({'category_code':code, 'brand':brand, 'products':count} for brand,count in sorted(Counter(r['brand'] for r in group).items()))
        plan.update(status='VERIFIED_SUBSET' if group else 'NO_VERIFIED_LISTING', review_status='VERIFIED_SUBSET' if group else 'NO_VERIFIED_LISTING', processing_rule='Include exact India store product listings with decoded local photos and source evidence; no stock check or exhaustive coverage claim.')
    write_csv(folder/'master_summary.csv', list(summaries[0]), summaries)
    write_csv(folder/'brand_summary.csv', ['category_code','brand','products'], brands)
    write_csv(folder/'category_coverage.csv', list(coverage[0]), coverage)
    write_csv(folder/'category_plan.csv', list(plans[0]), plans)
    write_csv(folder/'release_year_summary.csv', ['category_code','release_year','products'], [{'category_code':p['department_code'],'release_year':'Unknown','products':sum(r['category_code']==p['department_code'] for r in rows)} for p in plans])
    sheet=wb.create_sheet('Summary');sheet.append(list(summaries[0]))
    for row in summaries:sheet.append(list(row.values()))
    wb.save(folder/'master_catalogue.xlsx')
    inventory=Workbook();sheet=inventory.active;sheet.title='Current India products';sheet.append(fields)
    for row in rows:sheet.append([row.get(key,'') for key in fields])
    inventory.save(ROOT/'product_inventory.xlsx')


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--staged',type=Path,default=ROOT/'.firecrawl'/'india-catalog-staged')
    parser.add_argument('--apply',action='store_true')
    args=parser.parse_args(); staged=args.staged.resolve(); folder=ROOT/'product_catalog'
    with (staged/'all_products.csv').open(encoding='utf-8-sig',newline='') as stream:
        reader=csv.DictReader(stream);rows=list(reader);fields=reader.fieldnames
    validate_snapshot(rows,staged/'images')
    evidence=json.loads((staged/'availability_evidence.json').read_text(encoding='utf-8'))
    if len(rows)<300 or len({r['brand'] for r in rows})<8 or len(rows)!=len(evidence):
        raise ValueError('Suspiciously small or incomplete replacement; inspect the staged report')
    by_name={e['product_name']:e for e in evidence}
    for row in rows:
        proof=by_name[row['product_name']]
        image=staged/'images'/row['image_url'].rsplit('/',1)[1]
        if proof['purchase_url']!=row['product_url'] or proof['price']!=row['official_price'] or proof['image_sha256']!=hashlib.sha256(image.read_bytes()).hexdigest():
            raise ValueError('Source evidence does not match product: '+row['product_name'])
    print(json.dumps({'mode':'apply' if args.apply else 'dry-run','products':len(rows),'brands':len({r['brand'] for r in rows}),'local_photos':len({r['image_url'] for r in rows})}),flush=True)
    if not args.apply:return
    backup=ROOT/'backend'/'recovery'/('catalog-replacement-'+datetime.now().strftime('%Y%m%dT%H%M%S'))
    backup.mkdir(parents=True)
    shutil.copytree(folder,backup/'product_catalog')
    shutil.copy2(ROOT/'product_inventory.xlsx',backup/'product_inventory.xlsx')
    with connect() as conn,conn.cursor() as cursor:
        for table in ('products','categories','product_assets','product_filter_values','product_identities','product_source_records','product_specifications','dealer_products','promotion_products','product_import_sources'):
            cursor.execute('SELECT * FROM '+table)
            with gzip.open(backup/(table+'.json.gz'),'wt',encoding='utf-8') as stream:json.dump(cursor.fetchall(),stream,default=str)
    with (folder/'category_plan.csv').open(encoding='utf-8-sig') as stream:plans=list(csv.DictReader(stream))
    exports(folder,rows,fields,plans)
    (folder/'images').mkdir(exist_ok=True)
    for row in rows:
        name=row['image_url'].rsplit('/',1)[1]
        shutil.copy2(staged/'images'/name,folder/'images'/name)
    for name in ('availability_evidence.json','verification_report.json'):
        shutil.copy2(staged/name,folder/name)
    # Old review leads and coverage claims are not part of the replacement source.
    for name in ('manufacturer_gap_additions.csv','manufacturer_gap_candidates.csv','manufacturer_gap_category_summary.csv','manufacturer_gap_review.csv','collection_log.csv','app_catalog_visibility_backup.csv','lg-ac-catalogue-2026.pdf'):
        target=(folder/name).resolve()
        if target.parent != folder.resolve():raise ValueError('Unsafe cleanup path')
        target.unlink(missing_ok=True)
    original_args=sys.argv
    try:
        sys.argv=['catalogue_app_import','--apply']
        catalogue_app_import.main()
    except Exception:
        # Importer rolls back its transaction. Restore the source snapshot as well.
        for source in (backup/'product_catalog').iterdir():
            if source.is_file():shutil.copy2(source,folder/source.name)
        shutil.copy2(backup/'product_inventory.xlsx',ROOT/'product_inventory.xlsx')
        raise
    finally:sys.argv=original_args
    with connect() as conn,conn.cursor() as cursor:
        hidden="p.catalog_visible=FALSE AND NOT EXISTS(SELECT 1 FROM bid_requests br WHERE br.product_id=p.id) AND NOT EXISTS(SELECT 1 FROM cart_items ci WHERE ci.product_id=p.id)"
        cursor.execute('SELECT COUNT(*) total FROM products p WHERE '+hidden);unreferenced=cursor.fetchone()['total']
        cursor.execute('DELETE sr FROM product_source_records sr JOIN products p ON p.id=sr.product_id WHERE '+hidden)
        cursor.execute('DELETE p FROM products p WHERE '+hidden)
        cursor.execute("UPDATE product_import_sources SET enabled=FALSE,next_run_at=NULL")
        conn.commit()
        cursor.execute('SELECT COUNT(*) total FROM products WHERE catalog_visible=FALSE');preserved=cursor.fetchone()['total']
        for table in ('bid_requests','cart_items'):
            cursor.execute(f'SELECT COUNT(*) total FROM {table} x LEFT JOIN product_reference_catalog p ON p.id=x.product_id WHERE p.id IS NULL')
            if cursor.fetchone()['total']:raise RuntimeError('Broken historical product references in '+table)
        cursor.execute("SELECT image_url FROM product_reference_catalog WHERE image_url LIKE '/catalog-images/%'")
        retained={r['image_url'].rsplit('/',1)[1] for r in cursor.fetchall()}
    for photo in (folder/'images').glob('*.webp'):
        if photo.resolve().parent!=(folder/'images').resolve():raise ValueError('Unsafe photo cleanup path')
        if photo.name not in retained:photo.unlink()
    result={'products':len(rows),'deleted_unreferenced_old_products':unreferenced,'history_products_preserved_hidden':preserved,'backup':str(backup.relative_to(ROOT))}
    (folder/'replacement_result.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
    documentation(folder,rows,json.loads((folder/'verification_report.json').read_text(encoding='utf-8')))
    print(json.dumps(result,indent=2),flush=True)


if __name__=='__main__':main()
